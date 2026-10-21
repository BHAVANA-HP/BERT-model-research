"""
Production Inference Script: Predict Medical vs Non-Medical on New Unseen URLs.
Incorporates:
1. Targeted HTML Extraction (focus on <article>/<main>, ignore <aside>/<nav>)
2. Top-K Chunk Averaging (average of top-3 highest chunks instead of solitary max)
3. Medical Chunk Density Filtering (suppresses isolated sidebar/tangential quotes on long pages)
"""

import os
import sys
import argparse
import asyncio
import csv
import json
import re
from pathlib import Path
from html.parser import HTMLParser
from typing import List, Dict, Tuple

os.environ["OMP_NUM_THREADS"] = "1"
os.environ["OPENBLAS_NUM_THREADS"] = "1"
os.environ["MKL_NUM_THREADS"] = "1"

import torch
from transformers import AutoTokenizer, AutoModelForSequenceClassification
import httpx
import openpyxl

torch.set_num_threads(4)
device = torch.device("cuda" if torch.cuda.is_available() else "cpu")

REPO_ROOT = Path(__file__).resolve().parent.parent
DEFAULT_MODEL_DIR = REPO_ROOT / "models" / "bert_medical_classifier"
FALLBACK_BASE_MODEL = "sentence-transformers/all-MiniLM-L6-v2"
CACHE_FILE = REPO_ROOT / "data" / "scraped_pages_cache.json"

CONCURRENCY = 30
TIMEOUT_SECONDS = 7.0
CHUNK_SIZE = 256
CHUNK_STRIDE = 64


# ============================================================
# 1. TARGETED HTML EXTRACTION
# ============================================================

class TargetedTextExtractor(HTMLParser):
    def __init__(self, has_main_container: bool = False):
        super().__init__()
        self.texts = []
        self.has_main = has_main_container
        self.in_main = not has_main_container
        self.ignore_depth = 0
        self.ignore_tags = {
            "script", "style", "noscript", "header", "footer", "nav", "svg",
            "aside", "select", "option", "iframe", "button", "form"
        }

    def handle_starttag(self, tag, attrs):
        tag_lower = tag.lower()
        if tag_lower in self.ignore_tags:
            self.ignore_depth += 1
        if self.has_main and tag_lower in ("article", "main"):
            self.in_main = True

    def handle_endtag(self, tag):
        tag_lower = tag.lower()
        if tag_lower in self.ignore_tags and self.ignore_depth > 0:
            self.ignore_depth -= 1
        if self.has_main and tag_lower in ("article", "main"):
            self.in_main = False

    def handle_data(self, data):
        if self.ignore_depth == 0 and self.in_main:
            cleaned = data.strip()
            if cleaned:
                self.texts.append(cleaned)

    def get_text(self):
        return " ".join(self.texts)


def extract_clean_text(html_content: str) -> str:
    try:
        # Check if page has structured article/main containers
        has_main = bool(re.search(r"<(article|main)\b", html_content, re.IGNORECASE))
        extractor = TargetedTextExtractor(has_main_container=has_main)
        extractor.feed(html_content)
        result = extractor.get_text()
        if len(result.split()) >= 30 or not has_main:
            return result
        # Fallback to general extraction if container was too sparse
        fallback = TargetedTextExtractor(has_main_container=False)
        fallback.feed(html_content)
        return fallback.get_text()
    except Exception:
        clean = re.sub(r"<(script|style|nav|footer|aside|header)[^>]*>.*?</\1>", " ", html_content, flags=re.DOTALL | re.IGNORECASE)
        clean = re.sub(r"<[^>]+>", " ", clean)
        return " ".join(clean.split())


async def fetch_page(client: httpx.AsyncClient, semaphore: asyncio.Semaphore, url_id: str, url: str) -> Tuple[str, str, str]:
    headers = {
        "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/122.0.0.0 Safari/537.36",
        "Accept": "text/html,application/xhtml+xml,application/xml;q=0.9,*/*;q=0.8",
    }
    fetch_url = url if url.startswith("http") else f"https://{url}"
    async with semaphore:
        try:
            response = await client.get(fetch_url, headers=headers, timeout=TIMEOUT_SECONDS, follow_redirects=True)
            if response.status_code == 200:
                clean_text = extract_clean_text(response.text)
                return url_id, clean_text, "success"
            return url_id, "", f"http_{response.status_code}"
        except httpx.TimeoutException:
            return url_id, "", "timeout"
        except Exception as e:
            return url_id, "", f"error_{type(e).__name__}"


def chunk_text(text: str, chunk_size: int = CHUNK_SIZE, stride: int = CHUNK_STRIDE) -> List[str]:
    words = text.split()
    if len(words) <= chunk_size:
        return [text] if text else []
    chunks = []
    for i in range(0, len(words), chunk_size - stride):
        chunk = " ".join(words[i:i + chunk_size])
        chunks.append(chunk)
        if i + chunk_size >= len(words):
            break
    return chunks


# ============================================================
# 2. INFERENCE WITH TOP-K AVERAGING & DENSITY FILTERING
# ============================================================

@torch.no_grad()
def score_page(model, tokenizer, text: str) -> Tuple[str, float]:
    if not text.strip():
        return "non_medical", 0.0

    chunks = chunk_text(text, chunk_size=CHUNK_SIZE, stride=CHUNK_STRIDE)
    if not chunks:
        return "non_medical", 0.0

    encodings = tokenizer(
        chunks,
        truncation=True,
        padding=True,
        max_length=256,
        return_tensors="pt"
    ).to(device)

    outputs = model(**encodings)
    probs = torch.softmax(outputs.logits, dim=1)[:, 1].cpu().tolist()

    # Improvement 1: Top-K Chunk Averaging (replaces single-chunk max)
    sorted_probs = sorted(probs, reverse=True)
    top_k = min(3, len(sorted_probs))
    agg_prob = sum(sorted_probs[:top_k]) / top_k

    # Improvement 2: Medical Chunk Density check for long pages (>= 5 chunks)
    # If < 10% of page chunks contain medical content, suppress isolated sidebar/quote
    if len(probs) >= 5:
        med_chunks = sum(1 for p in probs if p >= 0.40)
        density = med_chunks / len(probs)
        if density < 0.10:
            agg_prob = agg_prob * 0.40

    label = "medical" if agg_prob >= 0.50 else "non_medical"
    return label, round(agg_prob, 4)


# ============================================================
# 3. INPUT FILE LOADER
# ============================================================

def load_urls_from_file(file_path: Path) -> List[Dict]:
    items = []
    if file_path.suffix.lower() == ".xlsx":
        wb = openpyxl.load_workbook(file_path, data_only=True)
        sheet = wb.active
        headers = [str(cell.value).strip().lower() if cell.value is not None else "" for cell in sheet[1]]
        url_col_idx = None
        id_col_idx = None
        for idx, h in enumerate(headers):
            if "url" in h and "id" not in h:
                url_col_idx = idx + 1
            if h.replace("-", "_") in ("url_id", "id"):
                id_col_idx = idx + 1
        if url_col_idx is None:
            url_col_idx = 1

        for row_idx in range(2, sheet.max_row + 1):
            url_val = sheet.cell(row=row_idx, column=url_col_idx).value
            if not url_val:
                continue
            url_str = str(url_val).strip()
            url_id = str(sheet.cell(row=row_idx, column=id_col_idx).value) if id_col_idx else f"URL_{row_idx-1:04d}"
            items.append({"url_id": url_id, "url": url_str})
    else:
        with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
            reader = csv.DictReader(f)
            url_col = None
            id_col = None
            for col in reader.fieldnames:
                if "url" in col.lower() and "id" not in col.lower():
                    url_col = col
                if col.lower().replace("-", "_") in ("url_id", "id"):
                    id_col = col
            if not url_col:
                url_col = reader.fieldnames[0]

            for idx, row in enumerate(reader, 1):
                url_val = row.get(url_col, "").strip()
                if not url_val:
                    continue
                url_id = row.get(id_col, f"URL_{idx:04d}").strip()
                items.append({"url_id": url_id, "url": url_val})
    return items


# ============================================================
# 4. MAIN PIPELINE
# ============================================================

async def run_prediction_pipeline(input_file: Path, output_file: Path, model_dir: Path):
    print("=" * 65)
    print("PRODUCTION BERT PREDICTOR (WITH TOP-K & DENSITY FILTERING)")
    print("=" * 65)
    print(f"Input file:  {input_file}")
    print(f"Output file: {output_file}")
    print(f"Model path:  {model_dir}")

    items = load_urls_from_file(input_file)
    print(f"Loaded {len(items)} URLs from input file.")
    if not items:
        print("Error: No URLs found to process.")
        return

    # Load Model
    if model_dir.exists():
        print(f"\n[Model] Loading fine-tuned classifier from: {model_dir}")
        tokenizer = AutoTokenizer.from_pretrained(model_dir)
        model = AutoModelForSequenceClassification.from_pretrained(model_dir).to(device)
    else:
        print(f"\n[Warning] Fine-tuned model not found at {model_dir}. Using base model: {FALLBACK_BASE_MODEL}")
        tokenizer = AutoTokenizer.from_pretrained(FALLBACK_BASE_MODEL)
        model = AutoModelForSequenceClassification.from_pretrained(FALLBACK_BASE_MODEL, num_labels=2).to(device)
    model.eval()

    # Load or fetch scraped texts
    cache = {}
    if CACHE_FILE.exists():
        try:
            with open(CACHE_FILE, "r", encoding="utf-8") as f:
                cache = json.load(f)
            print(f"[Cache] Found {len(cache)} existing cached pages on disk.")
        except Exception:
            cache = {}

    to_fetch = [item for item in items if item["url_id"] not in cache]
    scraped_map = {uid: (cache[uid].get("text", ""), cache[uid].get("status", "cached")) for uid in cache}

    if to_fetch:
        print(f"\n[Scraper] Fetching {len(to_fetch)} new URLs (concurrency={CONCURRENCY})...")
        semaphore = asyncio.Semaphore(CONCURRENCY)
        async with httpx.AsyncClient(verify=False) as client:
            tasks = [fetch_page(client, semaphore, item["url_id"], item["url"]) for item in to_fetch]
            scraped_results = await asyncio.gather(*tasks)
        for uid, txt, stat in scraped_results:
            scraped_map[uid] = (txt, stat)

    # Predict
    print(f"\n[Inference] Scoring with Top-K Averaging & Density Filtering...")
    results = []
    medical_count = 0
    for idx, item in enumerate(items, 1):
        uid = item["url_id"]
        url = item["url"]
        text, status = scraped_map.get(uid, ("", "unknown"))
        word_count = len(text.split())

        pred_label, conf = score_page(model, tokenizer, text)
        if pred_label == "medical":
            medical_count += 1

        results.append({
            "url_id": uid,
            "url": url,
            "word_count": word_count,
            "predicted_label": pred_label,
            "confidence_score": conf,
            "status": status
        })

        if idx % 200 == 0 or idx == len(items):
            print(f"  Processed {idx}/{len(items)} URLs (Medical calls so far: {medical_count})")

    # Save output Excel & CSV
    output_csv = output_file.with_suffix(".csv")
    with open(output_csv, "w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=["url_id", "url", "word_count", "predicted_label", "confidence_score", "status"])
        writer.writeheader()
        writer.writerows(results)

    output_xlsx = output_file.with_suffix(".xlsx")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "BERT Predictions"
    headers = ["url-id", "url", "word_count", "predicted_label", "confidence_score", "status"]
    ws.append(headers)
    for r in results:
        ws.append([r["url_id"], r["url"], r["word_count"], r["predicted_label"], r["confidence_score"], r["status"]])
    wb.save(output_xlsx)

    print("\n" + "=" * 65)
    print("PREDICTION COMPLETE")
    print("=" * 65)
    print(f"Total Processed: {len(results)}")
    print(f"Medical:         {medical_count} ({medical_count / len(results) * 100:.2f}%)")
    print(f"Non-Medical:     {len(results) - medical_count}")
    print(f"CSV Saved to:    {output_csv}")
    print(f"Excel Saved to:  {output_xlsx}")
    print("=" * 65)


def main():
    parser = argparse.ArgumentParser(description="Predict Medical vs Non-Medical on new URLs using fine-tuned BERT")
    parser.add_argument("--input", "-i", type=Path, required=True, help="Path to input .xlsx or .csv containing URLs")
    parser.add_argument("--output", "-o", type=Path, default=Path("results/new_urls_predictions"), help="Output path (without extension)")
    parser.add_argument("--model", "-m", type=Path, default=DEFAULT_MODEL_DIR, help="Path to fine-tuned model directory")
    args = parser.parse_args()

    asyncio.run(run_prediction_pipeline(args.input, args.output, args.model))


if __name__ == "__main__":
    main()
