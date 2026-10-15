import os
import asyncio
import csv
import math
import re
import sys
import time
from html.parser import HTMLParser
from pathlib import Path
import httpx
import openpyxl
import torch
from sentence_transformers import SentenceTransformer, util

torch.set_num_threads(4)

DOWNLOADS_DIR = Path.home() / "Downloads"

def find_file(pattern: str, fallback: str) -> Path:
    if DOWNLOADS_DIR.exists():
        matches = list(DOWNLOADS_DIR.glob(pattern))
        if matches:
            return matches[0]
    return Path(fallback)

FILE_PATH = Path(sys.argv[1]) if len(sys.argv) > 1 else (
    Path(os.environ["URL_EXCEL_PATH"]) if "URL_EXCEL_PATH" in os.environ else find_file("*classified_reference*.xlsx", "data/classified_reference.xlsx")
)
CSV_PATH = Path(sys.argv[2]) if len(sys.argv) > 2 else (
    Path(os.environ["URL_CSV_PATH"]) if "URL_CSV_PATH" in os.environ else find_file("*classified_reference*.csv", "data/classified_reference.csv")
)
CHECKPOINT_PATH = Path(os.environ.get("CHECKPOINT_PATH", DOWNLOADS_DIR / "checkpoint_bert_eval.csv" if DOWNLOADS_DIR.exists() else "checkpoint_bert_eval.csv"))

MODEL_NAME = "sentence-transformers/all-MiniLM-L6-v2"
HEALTH_DESC = "a page about medical diagnosis, disease symptoms, clinical treatment, prescription medicine, or healthcare"
NOT_HEALTH_DESC = "a page not about health or medicine"

CONCURRENCY = 30
TIMEOUT_SECONDS = 7.0


# ============================================================
# 1. HTML PARSER
# ============================================================

class TextExtractor(HTMLParser):
    def __init__(self):
        super().__init__()
        self.texts = []
        self.ignore = False
        self.ignore_tags = {
            "script", "style", "noscript", "header", "footer", "nav", "svg",
            "aside", "select", "option", "iframe", "button"
        }

    def handle_starttag(self, tag, attrs):
        if tag.lower() in self.ignore_tags:
            self.ignore = True

    def handle_endtag(self, tag):
        if tag.lower() in self.ignore_tags:
            self.ignore = False

    def handle_data(self, data):
        if not self.ignore:
            cleaned = data.strip()
            if cleaned:
                self.texts.append(cleaned)

    def get_text(self):
        return " ".join(self.texts)


def extract_clean_text(html_content: str) -> str:
    try:
        extractor = TextExtractor()
        extractor.feed(html_content)
        return extractor.get_text()
    except Exception:
        clean = re.sub(r"<(script|style|nav|footer)[^>]*>.*?</\1>", " ", html_content, flags=re.DOTALL | re.IGNORECASE)
        clean = re.sub(r"<[^>]+>", " ", clean)
        return " ".join(clean.split())


# ============================================================
# 2. ASYNC FETCH WORKER
# ============================================================

async def fetch_page(client: httpx.AsyncClient, semaphore: asyncio.Semaphore, url_id: str, url: str):
    headers = {
        "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/122.0.0.0 Safari/537.36",
        "Accept": "text/html,application/xhtml+xml,application/xml;q=0.9,*/*;q=0.8",
    }
    fetch_url = url if url.startswith("http") else f"https://{url}"

    async with semaphore:
        try:
            response = await client.get(fetch_url, headers=headers, timeout=TIMEOUT_SECONDS, follow_redirects=True)
            if response.status_code == 200:
                clean_text = extract_clean_text(response.text)
                return url_id, clean_text, "success"
            return url_id, "", f"http_{response.status_code}"
        except httpx.TimeoutException:
            return url_id, "", "timeout"
        except Exception as e:
            return url_id, "", f"error_{type(e).__name__}"


# ============================================================
# 3. MAIN EVALUATION PIPELINE
# ============================================================

async def main():
    all_rows = []
    if FILE_PATH.exists():
        print(f"Loading reference sheet from Excel: {FILE_PATH} ...")
        wb = openpyxl.load_workbook(FILE_PATH)
        ws = wb.active
        headers = [cell.value for cell in ws[1]]
        for row in ws.iter_rows(min_row=2, values_only=True):
            all_rows.append(dict(zip(headers, row)))
    elif CSV_PATH.exists():
        print(f"Loading reference sheet from CSV: {CSV_PATH} ...")
        with CSV_PATH.open("r", encoding="utf-8") as f:
            reader = csv.DictReader(f)
            for row in reader:
                all_rows.append(row)
    else:
        raise FileNotFoundError(f"Neither {FILE_PATH} nor {CSV_PATH} was found!")

    print(f"Loaded {len(all_rows)} reference rows.")

    # Blinding: Extract ONLY url-id and url
    blinded_inputs = [{"url-id": str(r["url-id"]), "url": str(r["url"])} for r in all_rows]

    # Load existing predictions from checkpoint if available
    bert_results = {}
    if CHECKPOINT_PATH.exists():
        with CHECKPOINT_PATH.open("r", encoding="utf-8") as f:
            reader = csv.DictReader(f)
            for r in reader:
                bert_results[r["url-id"]] = {
                    "bert_confidence_score": float(r["bert_confidence_score"]),
                    "bert_label": r["bert_label"],
                }
        print(f"Loaded {len(bert_results)} previous predictions from checkpoint.")

    remaining_inputs = [b for b in blinded_inputs if b["url-id"] not in bert_results]
    print(f"URLs remaining to process with BERT: {len(remaining_inputs)}")

    if remaining_inputs:
        print(f"Loading BERT model: {MODEL_NAME} ...")
        t_model = time.perf_counter()
        model = SentenceTransformer(MODEL_NAME, device="cpu")
        print(f"Model loaded in {time.perf_counter() - t_model:.2f} seconds.")

        # Embed anchor descriptions
        desc_embs = model.encode([HEALTH_DESC, NOT_HEALTH_DESC], normalize_embeddings=True, convert_to_tensor=True)
        health_emb = desc_embs[0]
        not_health_emb = desc_embs[1]

        semaphore = asyncio.Semaphore(CONCURRENCY)
        limits = httpx.Limits(max_keepalive_connections=CONCURRENCY, max_connections=CONCURRENCY * 2)

        cp_exists = CHECKPOINT_PATH.exists() and CHECKPOINT_PATH.stat().st_size > 0
        cp_file = CHECKPOINT_PATH.open("a" if cp_exists else "w", newline="", encoding="utf-8")
        cp_writer = csv.DictWriter(cp_file, fieldnames=["url-id", "url", "bert_confidence_score", "bert_label"])
        if not cp_exists:
            cp_writer.writeheader()
            cp_file.flush()

        batch_size = 50
        total_batches = math.ceil(len(remaining_inputs) / batch_size)
        t_start = time.perf_counter()
        processed_count = 0

        async with httpx.AsyncClient(limits=limits, verify=False) as client:
            for b_idx in range(total_batches):
                batch = remaining_inputs[b_idx * batch_size:(b_idx + 1) * batch_size]
                fetch_tasks = [fetch_page(client, semaphore, item["url-id"], item["url"]) for item in batch]
                fetched_pages = await asyncio.gather(*fetch_tasks)

                texts_to_encode = []
                valid_indices = []
                batch_records = []

                for uid, text, status in fetched_pages:
                    words = text.split()
                    if len(words) >= 15:
                        sample = " ".join(words[:250])
                        texts_to_encode.append(sample)
                        valid_indices.append(uid)
                    else:
                        batch_records.append((uid, 0.0, "non_medical"))

                if texts_to_encode:
                    embs = model.encode(texts_to_encode, batch_size=32, normalize_embeddings=True, convert_to_tensor=True)
                    sim_h = util.cos_sim(embs, health_emb).squeeze(1)
                    sim_nh = util.cos_sim(embs, not_health_emb).squeeze(1)

                    for idx, uid in enumerate(valid_indices):
                        sh = sim_h[idx].item()
                        snh = sim_nh[idx].item()
                        margin = sh - snh
                        prob = 1.0 / (1.0 + math.exp(-15.0 * margin))
                        prob = round(min(max(prob, 0.0), 1.0), 4)
                        label = "medical" if margin > 0.0 else "non_medical"
                        batch_records.append((uid, prob, label))

                batch_url_map = {item["url-id"]: item["url"] for item in batch}
                for uid, score, label in batch_records:
                    bert_results[uid] = {
                        "bert_confidence_score": score,
                        "bert_label": label,
                    }
                    cp_writer.writerow({
                        "url-id": uid,
                        "url": batch_url_map[uid],
                        "bert_confidence_score": score,
                        "bert_label": label,
                    })

                cp_file.flush()
                processed_count += len(batch)
                elapsed = time.perf_counter() - t_start
                rate = processed_count / elapsed if elapsed > 0 else 0
                eta = (len(remaining_inputs) - processed_count) / rate if rate > 0 else 0

                bert_med_count = sum(1 for r in bert_results.values() if r["bert_label"] == "medical")
                print(
                    f"[{processed_count}/{len(remaining_inputs)}] ({processed_count/len(remaining_inputs)*100:.1f}%) | "
                    f"Rate: {rate:.1f} urls/s | ETA: {eta:.0f}s | "
                    f"BERT Medical: {bert_med_count}"
                )

        cp_file.close()

    # ============================================================
    # 4. UNBLIND, MERGE & EVALUATE
    # ============================================================
    print("\nUnblinding reference ground truth and compiling 11 columns...")

    final_columns = [
        "url-id", "url", "word_count", "label", "confidence_score",
        "evidence_terms", "status", "bert_confidence_score", "bert_label",
        "agreement", "disagreement_type"
    ]

    updated_rows = []
    dict_rows = []
    tp = 0
    fp = 0
    fn = 0
    tn = 0

    for r in all_rows:
        uid = str(r["url-id"])
        ref_label = r["label"]
        pred_info = bert_results.get(uid, {"bert_confidence_score": 0.0, "bert_label": "non_medical"})

        bert_score = pred_info["bert_confidence_score"]
        bert_label = pred_info["bert_label"]

        if bert_label == ref_label:
            agreement = "MATCH"
            disagree_type = "-"
            if ref_label == "medical":
                tp += 1
            else:
                tn += 1
        else:
            agreement = "MISMATCH"
            if ref_label == "non_medical" and bert_label == "medical":
                disagree_type = "False Positive (BERT extra)"
                fp += 1
            else:
                disagree_type = "False Negative (BERT missed)"
                fn += 1

        w_cnt = int(r.get("word_count") or 0)
        c_score = float(r.get("confidence_score") or 0.0)

        row_vals = [
            uid,
            r["url"],
            w_cnt,
            ref_label,
            c_score,
            r.get("evidence_terms") or "",
            r.get("status") or "success",
            bert_score,
            bert_label,
            agreement,
            disagree_type,
        ]
        updated_rows.append(row_vals)
        dict_rows.append(dict(zip(final_columns, row_vals)))

    print(f"\nSaving updated unified Excel to {FILE_PATH} ...")
    wb_out = openpyxl.Workbook()
    ws_out = wb_out.active
    ws_out.title = "Classified Reference & BERT"
    ws_out.append(final_columns)

    for row in updated_rows:
        ws_out.append(row)

    wb_out.save(FILE_PATH)
    print(f"Successfully saved all 11 columns to Excel: {FILE_PATH}!")

    print(f"Saving updated unified CSV to {CSV_PATH} ...")
    with CSV_PATH.open("w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=final_columns)
        writer.writeheader()
        writer.writerows(dict_rows)
    print(f"Successfully saved all 11 columns to CSV: {CSV_PATH}!")

    total = len(updated_rows)
    accuracy = round((tp + tn) / total, 4) if total > 0 else 0.0
    precision = round(tp / (tp + fp), 4) if (tp + fp) > 0 else 0.0
    recall = round(tp / (tp + fn), 4) if (tp + fn) > 0 else 0.0
    f1 = round(2 * precision * recall / (precision + recall), 4) if (precision + recall) > 0 else 0.0
    agreement_rate = round((tp + tn) / total * 100, 2) if total > 0 else 0.0

    print("\n" + "=" * 65)
    print("BERT MODEL EVALUATION AGAINST REFERENCE GROUND TRUTH")
    print("=" * 65)
    print(f"Total URLs Evaluated:       {total}")
    print(f"Reference Medical Count:    {tp + fn}")
    print(f"BERT Medical Count:         {tp + fp}")
    print(f"Agreement Rate:             {agreement_rate}% ({tp + tn}/{total} matches)")
    print("-" * 65)
    print(f"Accuracy:                   {accuracy:.4f}")
    print(f"Precision:                  {precision:.4f}")
    print(f"Recall:                     {recall:.4f}")
    print(f"F1-Score:                   {f1:.4f}")
    print("-" * 65)
    print("CONFUSION MATRIX:")
    print(f"  True Positives (TP):      {tp} (Both agreed: Medical)")
    print(f"  True Negatives (TN):      {tn} (Both agreed: Non-Medical)")
    print(f"  False Positives (FP):     {fp} (BERT predicted medical, Ref was non-med)")
    print(f"  False Negatives (FN):     {fn} (BERT missed medical page)")
    print("=" * 65)


if __name__ == "__main__":
    asyncio.run(main())
